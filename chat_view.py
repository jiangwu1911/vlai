import logging
import sys
from PyQt6.QtWidgets import (QApplication, QDialog, QWidget, QVBoxLayout,
                            QHBoxLayout, QListWidget, QListWidgetItem, QLabel, 
                            QTextBrowser, QSplitter, QMessageBox, QPushButton,
                            QLineEdit)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QPixmap, QFont
from PyQt6.QtWidgets import QFileDialog
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
import os
import markdown
import json
from chat_model import ChatSession, ChatDatabase, ChatContent
from chat_utils import *
from multi_lang import LANGUAGES

db = ChatDatabase()

class ChatViewer(QDialog):
    def __init__(self, db_path='chat_sessions.db', parent=None, language='zh'):
        super().__init__(parent)
        self.language = language
        self.db_path = db_path
        self.init_database()

        self.setWindowTitle(LANGUAGES[self.language]["chat_viewer"])
        self.setGeometry(100, 100, 1000, 700)
        self.main_layout = QVBoxLayout(self)
        self.create_search_box()
        splitter_layout = QHBoxLayout()
        self.main_layout.addLayout(splitter_layout)
        self.splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter_layout.addWidget(self.splitter)

        self.create_session_list_widget()
        self.create_detail_panel()

        self.load_sessions()
        self.create_buttons()
        self.auto_select_first_session()

    def init_database(self):
        """初始化数据库连接"""
        if not os.path.exists(self.db_path):
            #QMessageBox.warning(self, "数据库不存在", f"找不到数据库文件: {self.db_path}")
            self.close()
            return

        self.engine = create_engine(f'sqlite:///{self.db_path}')
        Session = sessionmaker(bind=self.engine)
        self.db_session = Session()

    def create_search_box(self):
        """创建搜索框"""
        search_layout = QHBoxLayout()
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText(LANGUAGES[self.language]["chat_view_search"])
        self.search_input.textChanged.connect(self.search_sessions)
        search_layout.addWidget(self.search_input)
        self.main_layout.addLayout(search_layout)

    def create_session_list_widget(self):
        """创建会话列表组件"""
        self.session_list_widget = QListWidget()
        self.session_list_widget.setAlternatingRowColors(True)
        self.session_list_widget.setMinimumWidth(250)
        self.splitter.addWidget(self.session_list_widget)

        self.session_list_widget.itemClicked.connect(self.on_session_selected)
        self.session_list_widget.itemSelectionChanged.connect(self.on_session_selected)

    def create_detail_panel(self):
        """创建详情面板"""
        detail_widget = QWidget()
        detail_layout = QVBoxLayout(detail_widget)
        self.splitter.addWidget(detail_widget)

        # 对话内容
        self.conversation_text = QTextBrowser()
        self.conversation_text.setFont(QFont("SimHei", 10))
        detail_layout.addWidget(self.conversation_text)

        # 设置分割器比例
        self.splitter.setSizes([250, 750])

    def load_sessions(self):
        """加载所有会话到列表"""
        try:
            sessions = self.db_session.query(ChatSession).order_by(ChatSession.start_time.desc()).all()
            if not sessions:
                return

            for session in sessions:
                item = QListWidgetItem()
                item.setText(session.start_time.strftime("%Y-%m-%d %H:%M:%S"))
                item.setData(Qt.ItemDataRole.UserRole, session.id)
                item.setCheckState(Qt.CheckState.Unchecked)  # 设置复选框状态
                self.session_list_widget.addItem(item)

        except Exception as e:
            logging.error(self, LANGUAGES[self.language]["error"],
                          f"{LANGUAGES[self.language]['chat_view_load_session_error']}: {(e)}")

    def search_sessions(self, keyword):
        """根据关键词搜索会话"""
        self.session_list_widget.clear()
        # 查询包含关键词的聊天内容对应的会话ID
        content_query = self.db_session.query(ChatContent.session_id).filter(
            ChatContent.content.ilike(f"%{keyword}%")
        ).distinct()
        session_ids = [result[0] for result in content_query]

        # 根据会话ID查询会话
        sessions = self.db_session.query(ChatSession).filter(
            ChatSession.id.in_(session_ids)
        ).order_by(ChatSession.start_time.desc()).all()

        for session in sessions:
            item = QListWidgetItem()
            item.setText(session.start_time.strftime("%Y-%m-%d %H:%M:%S"))
            item.setData(Qt.ItemDataRole.UserRole, session.id)
            item.setCheckState(Qt.CheckState.Unchecked)
            self.session_list_widget.addItem(item)

    def remove_new_line(self, str):
        if str.startswith('<p>') and str.endswith('</p>'):
            str1 = str[3:-4]
            return str1
        return str

    def escape_response(self, str):
        str = str.replace("<p>", "")
        str = str.replace("</p>", "<br>")
        str = str.replace("<think>", "&lt;think&gt;<span style=\"color:gray\">")
        str = str.replace("</think>", "</span>&lt;/think&gt;")
        return str

    def on_session_selected(self, item=None):
        if item is None:
            selected_items = self.session_list_widget.selectedItems()
            if not selected_items:
                return
            item = selected_items[0]

        session_id = item.data(Qt.ItemDataRole.UserRole)
        try:
            # 使用 Session.get() 替代 Query.get()
            session = self.db_session.get(ChatSession, session_id)
            if not session:
                QMessageBox.warning(self, LANGUAGES[self.language]["chat_view_session_not_found"],
                                    LANGUAGES[self.language]["chat_view_session_deleted"])
                return

            # 显示对话内容
            chat_contents = db.search_chat_content_by_session_id(session_id)
            html_text = ""

            for content in chat_contents:
                if content.type == "Question":
                    question = self.remove_new_line(markdown.markdown(content.content))
                    user_html = f'<img src="{get_app_dir()}/assets/user.png" style="vertical-align: middle; ">&nbsp;&nbsp; {question}'
                    html_text += user_html

                if content.type == "Answer":
                    response = self.remove_new_line(markdown.markdown(content.content))
                    assistant_html = f'<br><img src="{get_app_dir()}/assets/robot.png" style="vertical-align: middle;"> &nbsp;&nbsp;&nbsp; {response}<br>'
                    html_text += f"<br>{assistant_html}<br>"

                if content.type == "Upload_image":
                    try:
                        images = json.loads(content.content)
                        if len(images) > 0:
                            table_html = "<table border='0' cellspacing='0' cellpadding='10'>"
                            count = 0
                            for image in images:
                                if count % 2 == 0:
                                    table_html += "<tr>"
                                count += 1
                                table_html += f"<td><img src='{image}' alt=LANGUAGES[self.language]['image'] width='320'></td>"
                            table_html += "</table><p>"
                            html_text += table_html
                    except Exception as e:
                        logging.error(f"Json parse error:{e}")

            self.conversation_text.setHtml(self.escape_response(html_text))

        except Exception as e:
            QMessageBox.critical(self, LANGUAGES[self.language]['error'],
                                 f"{LANGUAGES[self.language]['chat_view_load_image_failed']}: {str(e)}")

    def create_buttons(self):
        """创建关闭和删除按钮"""
        button_layout = QHBoxLayout()

        # 左侧按钮（关闭和删除）
        left_buttons = QHBoxLayout()
        close_button = QPushButton(LANGUAGES[self.language]['close'], self)
        close_button.clicked.connect(self.close)
        left_buttons.addWidget(close_button)

        delete_button = QPushButton(LANGUAGES[self.language]['chat_view_delete_selected_session'], self)
        delete_button.clicked.connect(self.delete_selected_sessions)
        left_buttons.addWidget(delete_button)

        # 将左侧按钮组添加到主布局
        button_layout.addLayout(left_buttons)

        # 添加弹性空间，将右侧按钮推到最右边
        button_layout.addStretch()

        # 右侧按钮（导出Word）
        export_word_button = QPushButton(LANGUAGES[self.language]['chat_view_export_to_word'], self)
        export_word_button.clicked.connect(self.export_to_word)
        button_layout.addWidget(export_word_button)

        self.main_layout.addLayout(button_layout)

    def delete_selected_sessions(self):
        """删除选中的会话"""
        selected_items = []
        for i in range(self.session_list_widget.count()):
            item = self.session_list_widget.item(i)
            if item.checkState() == Qt.CheckState.Checked:
                selected_items.append(item)

        if not selected_items:
            QMessageBox.information(self, LANGUAGES[self.language]['chat_view_no_selected'],
                                    LANGUAGES[self.language]['chat_view_choose_session'])
            return

        reply = QMessageBox.question(self, LANGUAGES[self.language]['chat_view_confirm'],
                                     LANGUAGES[self.language]['chat_view_confirm_delete_session'],
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No, QMessageBox.StandardButton.No)
        if reply == QMessageBox.StandardButton.Yes:
            for item in selected_items:
                session_id = item.data(Qt.ItemDataRole.UserRole)
                session = self.db_session.get(ChatSession, session_id)
                if session:
                    # 删除该会话对应的聊天内容
                    self.db_session.query(ChatContent).filter(
                        ChatContent.session_id == session_id
                    ).delete()
                    self.db_session.delete(session)
                self.session_list_widget.takeItem(self.session_list_widget.row(item))
            self.db_session.commit()

    def auto_select_first_session(self):
        if self.session_list_widget.count() > 0:
            first_item = self.session_list_widget.item(0)
            self.session_list_widget.setCurrentItem(first_item)
            self.on_session_selected(first_item)

    def export_to_word(self):
        """导出当前会话为Word文档"""
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        from datetime import datetime
        import os

        selected_items = self.session_list_widget.selectedItems()
        if not selected_items:
            QMessageBox.warning(self, LANGUAGES[self.language]['warning'],
                                LANGUAGES[self.language]["chat_view_select_a_session"])
            return

        session_id = selected_items[0].data(Qt.ItemDataRole.UserRole)
        session = self.db_session.get(ChatSession, session_id)
        if not session:
            QMessageBox.warning(self, LANGUAGES[self.language]['error'],
                                LANGUAGES[self.language]["chat_view_session_not_found"])
            return

        try:
            # 创建Word文档
            doc = Document()

            # 添加标题
            title = doc.add_paragraph()
            title_run = title.add_run(f"聊天记录 - {session.start_time.strftime('%Y-%m-%d %H:%M:%S')}")
            title_run.bold = True
            title_run.font.size = Pt(14)
            title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

            doc.add_paragraph()  # 空行

            # 获取聊天内容
            chat_contents = db.search_chat_content_by_session_id(session_id)

            for content in chat_contents:
                if content.type == "Question":
                    p = doc.add_paragraph()
                    p.add_run("用户: ").bold = True
                    p.add_run(content.content)

                elif content.type == "Answer":
                    p = doc.add_paragraph()
                    p.add_run("助手: ").bold = True
                    p.add_run(content.content)
                    doc.add_paragraph()  # 空行

                elif content.type == "Upload_image":
                    try:
                        images = json.loads(content.content)
                        for image in images:
                            if os.path.exists(image):
                                doc.add_paragraph(LANGUAGES[self.language]['image'])
                                doc.add_picture(image, width=Inches(4.0))
                                doc.add_paragraph()  # 空行
                    except Exception as e:
                        logging.error(f"导出图片失败: {e}")

            # 保存文件
            default_filename = f"chat_{session.start_time.strftime('%Y%m%d_%H%M%S')}.docx"
            save_path, _ = QFileDialog.getSaveFileName(
                self, LANGUAGES[self.language]['chat_view_save_word'], default_filename, "Word文档 (*.docx)"
            )

            if save_path:
                doc.save(save_path)
                QMessageBox.information(self, LANGUAGES[self.language]['success'],
                                        f"{LANGUAGES[self.language]['chat_view_save_word_to']}: {save_path}")

        except Exception as e:
            QMessageBox.critical(self, LANGUAGES[self.language]['error'],
                                       f"{LANGUAGES[self.language]['chat_view_save_word_failed']}")


if __name__ == "__main__":
    app = QApplication(sys.argv)

    # 确保中文显示正常
    font = QFont("SimHei")
    app.setFont(font)

    # 如果提供了数据库路径作为参数，则使用该路径
    db_path = 'history/chat_sessions.db'
    if len(sys.argv) > 1:
        db_path = sys.argv[1]

    viewer = ChatViewer(db_path)
    viewer.show()

    sys.exit(app.exec())
